"""CSV/Excel bill import: presets, parsing, duplicate checks and import jobs."""
from collections import Counter
from contextlib import closing, contextmanager
import csv
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from difflib import SequenceMatcher
import hashlib
import io
import json
import os
from pathlib import Path
import random
import re
import sqlite3
import time
from typing import Literal
import unicodedata
from uuid import uuid4
from zipfile import ZipFile, BadZipFile
from zoneinfo import ZoneInfo

from filelock import FileLock
from pydantic import BaseModel, ConfigDict, Field, model_validator
import regex
import requests
from service_auth import agent_requests

import assets
from categorization import category
from database import DBmanage
from transaction_reconciliation import (
    attach_source, ensure_reconciliation_schema, find_cross_source_match,
    merge_merchant_if_generic, source_id_for_import,
)
from runtime_settings import server_data_path


# Preset configuration

FIELDS = {"time", "amount", "shop_name", "product_name", "direction", "status",
          "transaction_type", "transaction_id", "merchant_order_id", "payment_method",
          "currency", "main_category", "sub_category"}
BUILTINS = Path(__file__).resolve().parent / "import_presets"


class ImportProblem(ValueError):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def data_root():
    return server_data_path("imports", "ABK_IMPORT_DATA_DIR")


def atomic_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name("." + path.name + "." + uuid4().hex + ".tmp")
    try:
        with temporary.open("w", encoding="utf-8", newline="\n") as stream:
            json.dump(data, stream, ensure_ascii=False, indent=2, allow_nan=False)
            stream.flush()
            os.fsync(stream.fileno())
        replace_file(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def replace_file(source, target, attempts=40):
    """os.replace that waits out readers: on Windows a file another thread has open (a status poll reading job.json)
    cannot be replaced until it is closed."""
    for attempt in range(attempts):
        try:
            return os.replace(source, target)
        except PermissionError:
            if attempt == attempts - 1:
                raise
            # Jitter keeps the retries from falling into step with a polling reader.
            time.sleep(0.002 + random.random() * 0.01 * (1 + attempt // 10))


class ReplaceRule(BaseModel):
    model_config = ConfigDict(extra="forbid")
    field: str
    pattern: str = Field(max_length=256)
    replacement: str = Field(default="", max_length=128)


class ImportPreset(BaseModel):
    model_config = ConfigDict(extra="forbid")
    preset_id: str = Field(pattern=r"^[a-z][a-z0-9_]{0,63}$")
    name: str = Field(min_length=1, max_length=100)
    provider: str = Field(pattern=r"^[a-z][a-z0-9_]{0,31}$")
    version: int = Field(ge=1)
    extensions: list[Literal[".csv", ".xlsx"]]
    encodings: list[Literal["utf-8-sig", "gb18030", "utf-16"]] = ["utf-8-sig", "gb18030"]
    default_currency: Literal["CNY", "JPY", "USD"] | None = None
    default_timezone: str = "Asia/Shanghai"
    columns: dict[str, str]
    expense_values: list[str] = Field(min_length=1, max_length=20)
    income_values: list[str] = Field(default_factory=list, max_length=20)
    success_statuses: list[str] = Field(default_factory=list, max_length=20)
    income_success_statuses: list[str] | None = Field(default=None, max_length=20)
    review_statuses: list[str] = Field(default_factory=list, max_length=20)
    excluded_type_patterns: list[str] = Field(default_factory=list, max_length=20)
    income_excluded_type_patterns: list[str] | None = Field(default=None, max_length=20)
    income_confirmation_patterns: list[str] = Field(default_factory=list, max_length=20)
    refund_patterns: list[str] = Field(default_factory=lambda: ["退款|退还|refund"], max_length=10)
    time_formats: list[str] = Field(default_factory=lambda: ["%Y-%m-%d %H:%M:%S"], max_length=10)
    replace_rules: list[ReplaceRule] = Field(default_factory=list, max_length=20)
    extract_patterns: dict[str, str] = Field(default_factory=dict)

    @model_validator(mode="after")
    def validate_rules(self):
        if not {"shop_name", "amount", "time", "direction"} <= self.columns.keys():
            raise ValueError("Map shop_name, amount, time and direction")
        if not set(self.columns) <= FIELDS or not set(self.extract_patterns) <= FIELDS:
            raise ValueError("Unsupported mapping field")
        if any(not value.strip() or len(value) > 100 for value in self.columns.values()):
            raise ValueError("Column names must contain 1..100 characters")
        if len(set(self.columns.values())) != len(self.columns):
            raise ValueError("Map each source column only once")
        if any(rule.field not in FIELDS for rule in self.replace_rules):
            raise ValueError("Unsupported replacement field")
        patterns = self.excluded_type_patterns + self.refund_patterns + list(self.extract_patterns.values())
        patterns += (self.income_excluded_type_patterns or []) + self.income_confirmation_patterns
        patterns += [rule.pattern for rule in self.replace_rules]
        for pattern in patterns:
            if len(pattern) > 256:
                raise ValueError("Regex is too long")
            regex.compile(pattern)
        if set(self.expense_values) & set(self.income_values):
            raise ValueError("Income and expense directions must be distinct")
        if any(len(value) > 100 for value in self.time_formats + self.expense_values + self.income_values
               + self.success_statuses + (self.income_success_statuses or []) + self.review_statuses):
            raise ValueError("Preset value is too long")
        ZoneInfo(self.default_timezone)
        return self


def regex_search(pattern, value):
    return regex.search(pattern, value, flags=regex.IGNORECASE, timeout=0.03)


class PresetStore:
    def __init__(self, root=None):
        self.root = Path(root or data_root()) / "presets"

    def load(self, preset_id):
        if not regex.fullmatch(r"[a-z][a-z0-9_]{0,63}", preset_id):
            raise ImportProblem("Invalid preset ID")
        for directory in (BUILTINS, self.root):
            path = directory / (preset_id + ".json")
            if path.is_file():
                return ImportPreset.model_validate_json(path.read_text(encoding="utf-8"))
        raise ImportProblem("Preset not found", 404)

    def all(self):
        presets = {}
        for directory in (self.root, BUILTINS):
            for path in directory.glob("*.json"):
                preset = ImportPreset.model_validate_json(path.read_text(encoding="utf-8"))
                presets[preset.preset_id] = preset
        return sorted(presets.values(), key=lambda item: item.preset_id)

    def save(self, preset):
        preset = ImportPreset.model_validate(preset)
        if (BUILTINS / (preset.preset_id + ".json")).exists():
            raise ImportProblem("Built-in presets are immutable; use a new preset ID", 409)
        self.root.mkdir(parents=True, exist_ok=True)
        path = self.root / (preset.preset_id + ".json")
        from filelock import FileLock
        with FileLock(str(self.root / ".presets.lock"), timeout=10):
            if path.exists():
                raise ImportProblem("Preset ID already exists; save a new version under a new ID", 409)
            atomic_json(path, preset.model_dump())
        return preset


# File parsing and normalization

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLUMNS = 100
MAX_CELL = 4000


def clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    text = unicodedata.normalize("NFKC", str(value)).strip()
    if len(text) > MAX_CELL:
        raise ImportProblem("Cell exceeds 4000 characters")
    return text


def read_tables(path, encodings=None):
    path = Path(path)
    if path.stat().st_size > MAX_BYTES:
        raise ImportProblem("File exceeds 10 MiB", 413)
    if path.suffix.lower() == ".csv":
        content = path.read_bytes()
        for encoding in encodings or ["utf-8-sig", "gb18030", "utf-16"]:
            try:
                text = content.decode(encoding, errors="strict")
                if "\x00" in text:
                    continue
                break
            except UnicodeError:
                continue
        else:
            raise ImportProblem("Unsupported CSV encoding")
        csv.field_size_limit(MAX_CELL * 4)
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows = []
        try:
            for row in reader:
                if len(rows) >= MAX_ROWS + 100 or len(row) > MAX_COLUMNS:
                    raise ImportProblem("File has too many rows or columns")
                rows.append([clean(value) for value in row])
        except csv.Error as error:
            raise ImportProblem("Invalid CSV: " + str(error)) from error
        return [{"name": "CSV", "rows": rows, "encoding": encoding}]
    if path.suffix.lower() != ".xlsx":
        raise ImportProblem("Only .csv and .xlsx are supported; convert legacy .xls first")
    try:
        with ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > 80 * 1024 * 1024:
                raise ImportProblem("Expanded workbook exceeds 80 MiB", 413)
            if len(archive.infolist()) > 10000:
                raise ImportProblem("Workbook contains too many archive entries")
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) > 20:
                raise ImportProblem("Workbook has too many sheets")
            tables = []
            for sheet in workbook.worksheets:
                if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + 100:
                    raise ImportProblem("Sheet has too many rows or columns")
                rows = []
                for row in sheet.iter_rows():
                    if len(rows) >= MAX_ROWS + 100:
                        raise ImportProblem("Sheet has too many rows")
                    # Formulas are preserved as text, never executed or trusted as cached money.
                    rows.append([clean(cell.value) for cell in row])
                tables.append({"name": sheet.title, "rows": rows, "encoding": None})
            return tables
        finally:
            workbook.close()
    except (BadZipFile, KeyError) as error:
        raise ImportProblem("Invalid XLSX workbook") from error


def locate_table(tables, preset, sheet_name=None):
    required = {clean(value) for value in preset.columns.values()}
    matches = []
    for table in tables:
        if sheet_name and table["name"] != sheet_name:
            continue
        for index, row in enumerate(table["rows"][:100]):
            headers = [clean(value) for value in row]
            if required <= set(headers):
                if len([h for h in headers if h]) != len(set(h for h in headers if h)):
                    raise ImportProblem("Duplicate column names")
                matches.append((table, index, headers))
                break
    if not matches:
        raise ImportProblem("Selected preset does not match the file headers", 422)
    if len(matches) > 1:
        raise ImportProblem("Multiple matching sheets; select sheet_name", 422)
    return matches[0]


def source_key(provider, transaction_id, file_hash, sheet, row_number):
    # A platform transaction ID survives overlapping exports; no amount-only dedupe.
    identity = [provider, transaction_id] if transaction_id else [provider, file_hash, sheet, row_number]
    import json
    return "file-import:" + hashlib.sha256(json.dumps(identity, ensure_ascii=True).encode()).hexdigest()


def parse_time(value, preset, timezone):
    if value.startswith("="):
        raise ValueError("Formula dates are not supported")
    try:
        result = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        result = None
        for format_string in preset.time_formats:
            try:
                result = datetime.strptime(value, format_string)
                break
            except ValueError:
                pass
        if result is None:
            raise ValueError("Unrecognized transaction time")
    return (result.replace(tzinfo=ZoneInfo(timezone)) if result.tzinfo is None else result).isoformat()


def normalize_rows(tables, preset, file_hash, account_number, timezone=None, currency=None, sheet_name=None):
    table, header_at, headers = locate_table(tables, preset, sheet_name)
    timezone = timezone or preset.default_timezone
    ZoneInfo(timezone)
    output = []
    for index, values in enumerate(table["rows"][header_at + 1:], header_at + 2):
        if not any(values):
            continue
        raw = {header: values[n] if n < len(values) else "" for n, header in enumerate(headers) if header}
        row = {"row_id": str(index), "sheet": table["name"], "row_number": index, "raw": raw,
               "normalized": {}, "status": "new", "issues": [], "duplicate_candidates": []}
        fields = {target: clean(raw.get(clean(column))) for target, column in preset.columns.items()}
        # Determine direction/status before transformations so rules cannot hide refunds.
        kind, direction, state = fields.get("transaction_type", ""), fields["direction"], fields.get("status", "")
        is_income = direction in preset.income_values
        excluded_patterns = (preset.income_excluded_type_patterns
            if is_income and preset.income_excluded_type_patterns is not None else preset.excluded_type_patterns)
        success_statuses = (preset.income_success_statuses
            if is_income and preset.income_success_statuses is not None else preset.success_statuses)
        try:
            if any(regex_search(pattern, kind + " " + state) for pattern in preset.refund_patterns):
                row.update(status="needs_review", issues=["refund_requires_review"])
            elif any(regex_search(pattern, kind) for pattern in excluded_patterns):
                row.update(status="excluded", issues=["excluded_transaction_type"])
            elif direction not in preset.expense_values and not is_income:
                row.update(status="excluded", issues=["not_an_expense"])
            elif state in preset.review_statuses:
                row.update(status="needs_review", issues=["payment_requires_review"])
            elif success_statuses and state not in success_statuses:
                row.update(status="excluded", issues=["not_a_successful_payment"])
            row["income_confirmation_required"] = is_income and any(regex_search(pattern,
                kind + " " + fields.get("product_name", "")) for pattern in preset.income_confirmation_patterns)
            if row["income_confirmation_required"]:
                row["issues"].append("confirm_external_income_required")
            for rule in preset.replace_rules:
                fields[rule.field] = regex.sub(rule.pattern, rule.replacement, fields.get(rule.field, ""), timeout=0.03)
            for field, pattern in preset.extract_patterns.items():
                match = regex_search(pattern, fields.get(field, ""))
                if not match:
                    raise ValueError("Extraction failed: " + field)
                fields[field] = match.group(1) if match.lastindex else match.group(0)
            source_currency = fields.get("currency")
            code = assets.currency(source_currency or preset.default_currency or currency or "")
            amount_text = regex.sub(r"^(?:CNY|RMB|JPY|USD|US\$|[¥￥$＄])\s*|\s*(?:元|円|ドル|美元|CNY|JPY|RMB|USD)$", "", fields["amount"], flags=regex.IGNORECASE).replace(",", "").strip()
            if not regex.fullmatch(r"-?\d+(?:\.\d+)?", amount_text):
                raise ValueError("Invalid amount")
            minor = assets.minor_amount(amount_text, code)
            if minor <= 0:
                raise ValueError("Transaction amount must be positive")
            if not fields["shop_name"] or fields["shop_name"] in {"/", "-"}:
                raise ValueError("Merchant name is missing")
            when = parse_time(fields["time"], preset, timezone)
            transaction_id = fields.get("transaction_id", "").strip()
            if transaction_id in {"/", "-"}:
                transaction_id = ""
            if transaction_id and (transaction_id.startswith("=") or regex.search(r"\d[eE][+-]\d", transaction_id)):
                raise ValueError("Transaction ID was converted to a formula or scientific notation")
            row["normalized"] = {"shop_name": fields["shop_name"], "raw_shop_name": fields["shop_name"],
                "product_name": fields.get("product_name", ""), "price": assets.major_amount(minor, code),
                "CurrencyTypeNumber": assets.currency_number(code), "isotime": when,
                "AccountNumber": account_number, "main_category": "", "sub_category": "", "detail_category": "",
                "payment_method": fields.get("payment_method", ""), "transaction_id": transaction_id,
                "merchant_order_id": fields.get("merchant_order_id", ""), "provider": preset.provider,
                "transaction_kind": "income" if is_income else "expense", "transaction_type": kind}
            row["operation_id"] = source_key(preset.provider, transaction_id, file_hash, table["name"], index)
        except (ValueError, TimeoutError) as error:
            row["issues"].append(str(error))
            if row["status"] == "new":
                row["status"] = "invalid"
        output.append(row)
    return output, {"sheet": table["name"], "header_row": header_at + 1, "headers": headers,
                    "encoding": table["encoding"], "timezone": timezone}


# Duplicate checks

def normalize_name(value):
    return re.sub(r"[\W_]+", "", unicodedata.normalize("NFKC", str(value or "")).casefold())


def similarity(left, right):
    left, right = normalize_name(left), normalize_name(right)
    if not left or not right or left in {"商品", "消费", "付款"} or right in {"商品", "消费", "付款"}:
        return 0.0
    if min(len(left), len(right)) <= 2:
        return 1.0 if left == right else 0.0
    return SequenceMatcher(None, left, right).ratio()


def find_duplicate_transaction(conn, row, window_minutes=1):
    if row["normalized"].get("transaction_kind") == "income":
        return find_duplicate_income(conn, row, window_minutes=window_minutes)
    data = row["normalized"]
    cross_source = find_cross_source_match(
        conn,
        amount=data["price"],
        currency_number=data["CurrencyTypeNumber"],
        isotime=data["isotime"],
        merchant=data.get("shop_name") or data.get("raw_shop_name"),
        source_type="statement",
        window_minutes=10,
    )
    if cross_source:
        cross_source["reasons"].append("cross_source_match")
        return "duplicate", [cross_source]
    op = conn.execute("SELECT ResultJSON FROM AssetOperation WHERE OperationID=?", (row["operation_id"],)).fetchone()
    if op:
        result = json.loads(op[0])
        number = result.get("DealingNumber")
        existing = assets.read_dealing(conn, number) if number is not None else None
        if existing is None:
            return "needs_review", [{"DealingNumber": number, "reasons": ["previous_import_was_deleted"]}]
        compatible = (assets.currency(existing["CurrencyTypeNumber"]) == assets.currency(data["CurrencyTypeNumber"])
                      and Decimal(str(existing["price"])) == Decimal(data["price"])
                      and abs((datetime.fromisoformat(assets.timestamp(existing["ISOtime"])) -
                               datetime.fromisoformat(assets.timestamp(data["isotime"]))).total_seconds()) <= 60)
        # The same row booked from this file into another account (a file split by payment method) is already in the
        # ledger: skip it and say where it is. Only a changed amount or time needs the user.
        reasons = ["same_import_identity"] if compatible else ["import_identity_conflict"]
        if compatible and str(existing.get("AccountNumber")) != str(data["AccountNumber"]):
            reasons.append("imported_to_other_account")
        return ("duplicate" if compatible else "needs_review"), [{
            "DealingNumber": number, "shop_name": existing["ShopName"], "price": existing["price"],
            "time": existing["ISOtime"], "AccountNumber": existing.get("AccountNumber"), "reasons": reasons}]
    when = datetime.fromisoformat(assets.timestamp(data["isotime"]))
    window = timedelta(minutes=window_minutes)
    # Compare endpoints to avoid floating-point subtraction at the 60-second boundary.
    cursor = conn.execute("""SELECT d.DealingNumber, d.ShopName, d.price, d.ISOtime, d.AccountNumber,
        (SELECT group_concat(a.ProductName, ' ') FROM ADRresults a WHERE a.DealingNumber=d.DealingNumber) AS Products
        FROM DealingData d WHERE d.CurrencyTypeNumber IN (?,?,?) AND ABS(d.price-?)<0.000001
        AND julianday(d.ISOtime) BETWEEN julianday(?) AND julianday(?)
        ORDER BY ABS(julianday(d.ISOtime)-julianday(?)), d.DealingNumber""",
        ((*{"CNY": (100, 0, 1), "JPY": (101, 101, 101), "USD": (102, 102, 102)}[assets.currency(data["CurrencyTypeNumber"])],
          data["price"], (when - window).isoformat(), (when + window).isoformat(), data["isotime"])))
    aliases = {data["shop_name"], data.get("raw_shop_name", "")}
    for key, name in conn.execute("SELECT key, ShopName FROM ShopName"):
        if normalize_name(key) in {normalize_name(a) for a in aliases} or normalize_name(name) in {normalize_name(a) for a in aliases}:
            aliases.update([key or "", name or ""])
    candidates = []
    for number, name, price, when, account, products in cursor.fetchall():
        merchant = max(similarity(alias, name) for alias in aliases)
        product = similarity(data.get("product_name"), products)
        if merchant >= 0.78 or product >= 0.65:
            reasons = ["same_amount_and_currency", "near_transaction_time"]
            if merchant >= 0.78:
                reasons.append("similar_merchant")
            if product >= 0.65:
                reasons.append("similar_product")
            candidates.append({"DealingNumber": number, "shop_name": name, "price": price, "time": when,
                               "account_number": account, "product_name": products, "reasons": reasons,
                               "merchant_similarity": round(merchant, 3), "product_similarity": round(product, 3)})
    return ("suspected_duplicate" if candidates else "new"), candidates[:20]


def find_duplicate_income(conn, row, window_minutes=1):
    data = row["normalized"]
    code = assets.currency(data["CurrencyTypeNumber"])
    amount = assets.minor_amount(data["price"], code)
    op = conn.execute("SELECT ResultJSON FROM AssetOperation WHERE OperationID=?", (row["operation_id"],)).fetchone()
    if op:
        result = json.loads(op[0])
        income_id = result.get("income_id")
        existing = assets.read_income(conn, income_id) if income_id is not None else None
        if existing is None:
            return "needs_review", [{"income_id": income_id, "reasons": ["previous_import_was_deleted_or_changed_kind"]}]
        compatible = (existing["Currency"] == code and existing["Amount"] == amount
                      and abs((datetime.fromisoformat(existing["EffectiveAt"]) -
                               datetime.fromisoformat(assets.timestamp(data["isotime"]))).total_seconds()) <= 60)
        reasons = ["same_import_identity"] if compatible else ["import_identity_conflict"]
        if compatible and existing["AccountNumber"] != data["AccountNumber"]:
            reasons.append("imported_to_other_account")
        return ("duplicate" if compatible else "needs_review"), [{"income_id": income_id,
            "shop_name": existing["SourceName"], "price": assets.major_amount(existing["Amount"], existing["Currency"]),
            "time": existing["EffectiveAt"], "AccountNumber": existing["AccountNumber"], "reasons": reasons}]
    candidates = []
    when = datetime.fromisoformat(assets.timestamp(data["isotime"]))
    window = timedelta(minutes=window_minutes)
    rows = conn.execute("""SELECT IncomeID,SourceName,Amount,EffectiveAt,Note FROM IncomeData
        WHERE AccountNumber=? AND Currency=? AND Amount=?
        AND julianday(EffectiveAt) BETWEEN julianday(?) AND julianday(?)
        ORDER BY ABS(julianday(EffectiveAt)-julianday(?)),IncomeID""",
        (data["AccountNumber"], code, amount, (when - window).isoformat(), (when + window).isoformat(), data["isotime"]))
    for income_id, name, value, when, note in rows:
        names = (data["shop_name"], data.get("raw_shop_name", ""))
        name_score = max(similarity(alias, name) for alias in names)
        note_score = similarity(data.get("product_name"), note)
        if not name or name_score >= 0.78 or note_score >= 0.65:
            candidates.append({"income_id": income_id, "shop_name": name, "price": assets.major_amount(value, code),
                "time": when, "account_number": data["AccountNumber"], "reasons": ["same_amount_and_currency",
                "near_transaction_time", "income_source_missing" if not name else "similar_income_source"]})
            if len(candidates) >= 20:
                break
    return ("suspected_duplicate" if candidates else "new"), candidates


# Import jobs and commits

class ImportService:
    SAVE_EVERY = 100          # commit: rows per ledger transaction and job-file save
    CLASSIFY_CHUNK = 200      # analysis: merchant/product pairs per Agent classify-batch request

    def __init__(self, db_path, agent_url, root=None, agent_client=None):
        self.db_path = str(Path(db_path).resolve())
        self.agent_url = agent_url.rstrip("/")
        self.root = Path(root or data_root())
        self.presets = PresetStore(self.root)
        self.agent_client = agent_client

    def folder(self, ident):
        if len(ident) != 32 or any(c not in "0123456789abcdef" for c in ident):
            raise ImportProblem("Invalid import ID")
        return self.root / "jobs" / ident

    def lock(self, ident):
        folder = self.folder(ident)
        if not (folder / "job.json").exists():
            raise ImportProblem("Import task not found", 404)
        return FileLock(str(folder / "job.lock"), timeout=1)

    def load(self, ident):
        try:
            return json.loads((self.folder(ident) / "job.json").read_text(encoding="utf-8"))
        except FileNotFoundError as error:
            raise ImportProblem("Import task not found", 404) from error

    def save(self, job):
        atomic_json(self.folder(job["import_id"]) / "job.json", job)

    @contextmanager
    def connect(self):
        conn = sqlite3.connect(self.db_path, timeout=15)
        conn.row_factory = sqlite3.Row
        try:
            ensure_reconciliation_schema(conn)
            with conn:
                yield conn
        finally:
            conn.close()

    def call_agent(self, route, payload):
        if self.agent_client:
            return self.agent_client(route, payload)
        response = agent_requests.post(self.agent_url + route, json=payload, timeout=(10, 120))
        response.raise_for_status()
        result = response.json()
        if not result.get("success"):
            raise ImportProblem(result.get("error") or "Agent request failed")
        return result["result"]

    def create(self, content, filename, preset_id, account_number, timezone_name=None,
               default_currency=None, sheet_name=None, use_agent=True):
        suffix = Path(filename).suffix.lower()
        if suffix not in {".csv", ".xlsx"}:
            raise ImportProblem("Only CSV and XLSX uploads are supported")
        if not content or len(content) > MAX_BYTES:
            raise ImportProblem("Empty file or file exceeds 10 MiB", 413)
        with self.connect() as conn:
            assets.get_account(conn, account_number)
        if preset_id:
            preset = self.presets.load(preset_id)
            if suffix not in preset.extensions:
                raise ImportProblem("File extension does not match selected preset")
        ident = uuid4().hex
        folder = self.folder(ident)
        folder.mkdir(parents=True)
        source = folder / ("source" + suffix)
        source.write_bytes(content)
        job = {"import_id": ident, "state": "queued", "revision": 0,
               "filename": Path(filename.replace("\\", "/")).name, "source_path": str(source),
               "file_hash": hashlib.sha256(content).hexdigest(), "preset_id": preset_id,
               "account_number": account_number, "timezone": timezone_name,
               "default_currency": default_currency, "sheet_name": sheet_name, "use_agent": use_agent,
               "rows": [], "warnings": [], "commits": {}, "created_at": datetime.now(timezone.utc).isoformat()}
        self.save(job)
        return self.summary(job)

    def suggest_mapping(self, job, tables):
        samples = []
        for table in tables:
            if job["sheet_name"] and table["name"] != job["sheet_name"]:
                continue
            populated = [(n, row) for n, row in enumerate(table["rows"][:100], 1) if sum(bool(x) for x in row) >= 4]
            # Limit private samples and prompt size; original files are not uploaded to the Agent.
            for number, row in populated[:5]:
                samples.append({"sheet": table["name"], "row_number": number, "cells": [str(x)[:150] for x in row[:20]]})
        if job["use_agent"] and samples:
            try:
                result = self.call_agent("/agent/import/preset-suggest", {"samples": samples, "extension": Path(job["source_path"]).suffix,
                                         "schema": ImportPreset.model_json_schema()})
                # The prompt asks for {"preset": {...}}; models also answer with the preset itself, which is just as
                # usable. An answer without columns is the model saying it could not map the file.
                preset = result.get("preset", result) if isinstance(result, dict) else None
                if not isinstance(preset, dict) or "columns" not in preset:
                    reason = result.get("error") if isinstance(result, dict) else None
                    raise ImportProblem(str(reason or "the model suggested no mapping")[:200])
                draft = ImportPreset.model_validate(preset)
                job["mapping_suggestion"] = draft.model_dump()
            except Exception as error:
                job["warnings"].append("preset_suggestion_failed: " + str(error)[:300])
        job["header_samples"] = samples
        job["state"] = "mapping_required"

    def prepare(self, ident, supplied_preset=None):
        with self.lock(ident):
            job = self.load(ident)
            if job["state"] not in {"queued", "mapping_required", "failed", "ready"} or job["commits"]:
                raise ImportProblem("This import can no longer be re-parsed", 409)
            job.update(state="parsing", error=None, rows=[], warnings=[])
            self.save(job)
            try:
                preset = ImportPreset.model_validate(supplied_preset) if supplied_preset else (
                    self.presets.load(job["preset_id"]) if job["preset_id"] else None)
                tables = read_tables(job["source_path"], preset.encodings if preset else None)
                if preset is None:
                    matching = []
                    for option in self.presets.all():
                        if Path(job["source_path"]).suffix not in option.extensions:
                            continue
                        try:
                            locate_table(tables, option, job["sheet_name"])
                            matching.append(option)
                        except ImportProblem:
                            pass
                    if len(matching) == 1:
                        preset = matching[0]
                if preset is None:
                    self.suggest_mapping(job, tables)
                else:
                    try:
                        rows, metadata = normalize_rows(tables, preset, job["file_hash"], job["account_number"],
                            job["timezone"], job["default_currency"], job["sheet_name"])
                    except ImportProblem as error:
                        if error.status != 422:
                            raise
                        job["warnings"].append(str(error))
                        self.suggest_mapping(job, tables)
                        self.save(job)
                        return
                    job.update(preset=preset.model_dump(), preset_id=preset.preset_id, table=metadata, state="analyzing", rows=rows)
                    self.save(job)
                    self.analyze(job)
                    job["state"] = "ready"
                    job["revision"] += 1
            except Exception as error:
                job.update(state="failed", error=str(error)[:500])
            self.save(job)

    def analyze(self, job):
        """Check, deduplicate and classify parsed rows in three phases.

        1. One ledger pass: currency, duplicates, same-file identity, the income rule, and merchants the ShopName
           table already knows.
        2. Every other merchant/product pair goes to the Agent's classify-batch in chunks, outside the ledger
           connection. The Agent answers from classification memory first and asks the model only for what memory
           does not know, merging an ordinary merchant's products into one question. Progress is saved per chunk.
        3. Apply the categories and re-check duplicates with the final merchant names.
        Every row is classified; nothing is cut off by a request budget. A row that fails is kept for review with its
        error and the rest continue. Answers already received survive an explicit re-parse in classification_cache."""
        manual = {"main_category": "", "sub_category": "", "source": "manual"}
        cache = {key: value for key, value in (job.get("classification_cache") or {}).items()
                 if value.get("main_category") and value.get("sub_category") and not value.get("error")}
        job["classification_cache"] = cache
        seen, known, targets = {}, {}, []
        with self.connect() as conn:
            account_currency = assets.currency(assets.get_account(conn, job["account_number"])["CurrencyType"])
            # ShopName can be empty when the existing classifier stores only vector memory.
            options = DBmanage(self.db_path).R_for_category_options(conn)[:500]
            for row in job["rows"]:
                try:
                    key = self._check_row(conn, row, seen, account_currency)
                    if key is None:
                        continue
                    targets.append((row, key))
                    if key in known:
                        continue
                    # The mail pipeline's lookup, read-only: an import preview never records unknown merchants.
                    shop, main, sub = category(row["normalized"]["CurrencyTypeNumber"]).Data_classification(
                        key[0], connection=conn, record_unknown=False)
                    if shop and main and sub and not {main, sub} & {"未分类", "待分类"}:
                        known[key] = {"shop_name": shop, "main_category": main, "sub_category": sub,
                                      "detail_category": "", "source": "shop_name_table"}
                    else:
                        known[key] = cache.get(self._cache_key(key))
                except Exception as error:
                    self._row_failed(row, error)
        pending = [key for key, value in known.items() if value is None]
        job["progress"] = {"stage": "classifying", "done": len(known) - len(pending), "total": len(known)}
        used = 0
        for start in range(0, len(pending), self.CLASSIFY_CHUNK):
            chunk = pending[start:start + self.CLASSIFY_CHUNK]
            asked = [key for key in chunk if key[0]]
            if job["use_agent"] and asked:
                try:
                    answer = self.call_agent("/agent/import/classify-batch", {
                        "items": [{"key": str(index), "shop_name": key[0], "product_name": key[1] or ""}
                                  for index, key in enumerate(asked)],
                        "category_options": options, "max_model_calls": 100})
                    used += int(answer.get("model_calls") or 0)
                    by_key = {str(result.get("key")): result for result in answer.get("results") or []}
                    for index, key in enumerate(asked):
                        known[key] = by_key.get(str(index)) or manual
                        if known[key].get("main_category") and known[key].get("sub_category") and not known[key].get("error"):
                            cache[self._cache_key(key)] = known[key]
                except Exception as error:
                    for key in asked:
                        known[key] = {"main_category": "", "sub_category": "", "source": "unavailable", "error": str(error)[:200]}
            for key in chunk:
                known[key] = known[key] or manual
            job["progress"]["done"] += len(chunk)
            self.save(job)
        with self.connect() as conn:
            for row, key in targets:
                try:
                    n, classification = row["normalized"], known[key]
                    for field in ("shop_name", "main_category", "sub_category", "detail_category"):
                        if classification.get(field):
                            n[field] = classification[field]
                    row["classification_source"] = classification.get("source")
                    if not n["main_category"] or not n["sub_category"]:
                        row["issues"].append("classification_required")
                    elif classification.get("needs_review"):
                        # The Agent picked a pair as the daily classifier does but was unsure; the row still imports.
                        row["issues"].append("classification_uncertain")
                    if classification.get("error"):
                        row["issues"].append(classification["error"])
                    if row["status"] == "new":
                        row["status"], row["duplicate_candidates"] = find_duplicate_transaction(conn, row)
                except Exception as error:
                    self._row_failed(row, error)
        job["agent_classification_calls"] = used
        job.pop("progress", None)

    @staticmethod
    def _check_row(conn, row, seen, account_currency):
        """Currency, duplicate, same-file identity and income checks for one parsed row. Returns the merchant/product
        key when the row still needs a category, otherwise None."""
        if row["status"] != "new":
            return None
        n = row["normalized"]
        if assets.currency(n["CurrencyTypeNumber"]) != account_currency:
            row.update(status="needs_review", issues=row["issues"] + ["account_currency_mismatch"])
            return None
        row["status"], row["duplicate_candidates"] = find_duplicate_transaction(conn, row)
        if row["operation_id"] in seen:
            earlier = seen[row["operation_id"]]
            if all(n.get(field) == earlier["normalized"].get(field) for field in
                   ("price", "CurrencyTypeNumber", "isotime", "AccountNumber", "raw_shop_name", "product_name", "transaction_kind")):
                row.update(status="duplicate", duplicate_candidates=[{"row_id": earlier["row_id"], "reasons": ["same_file_identity"]}])
            else:
                row.update(status="needs_review", issues=row["issues"] + ["conflicting_rows_with_same_identity"])
        seen[row["operation_id"]] = row
        if row["status"] not in {"new", "suspected_duplicate"}:
            return None
        if n.get("transaction_kind") == "income":
            n.update(main_category="收入", sub_category=n.get("transaction_type") or "其他收入")
            row["classification_source"] = "income_rule"
            return None
        return n["raw_shop_name"], n["product_name"]

    @staticmethod
    def _row_failed(row, error):
        row.update(status="needs_review", processing_error=str(error)[:500])
        row["issues"].append("row_processing_failed")

    @staticmethod
    def _cache_key(key):
        return json.dumps(list(key), ensure_ascii=False)

    @staticmethod
    def _open_expense(row):
        """An expense still to import, as (is it one, does it still lack a category)."""
        n = row["normalized"]
        expense = (row["status"] in {"new", "suspected_duplicate"} and n.get("transaction_kind") != "income"
                   and row.get("commit_status") != "imported")
        return expense, expense and not (n.get("main_category") and n.get("sub_category"))

    # What the review page shows: rows ready to import, rows that need a look, rows already in the ledger (skipped on
    # their own), rows that are not bookkeeping (transfers, fund income...), and rows already handled.
    GROUPS = ("ready", "attention", "existing", "excluded", "done")

    @classmethod
    def row_group(cls, row):
        state = row.get("commit_status") or row["status"]
        if state in {"imported", "skipped"}:
            return "done"
        if state == "duplicate":
            return "existing"
        if state == "excluded":
            return "excluded"
        income_check = row.get("income_confirmation_required") or "confirm_external_income_required" in (row.get("issues") or [])
        if state == "new" and not cls._open_expense(row)[1] and not income_check:
            return "ready"
        return "attention"

    @staticmethod
    def previous_accounts(job):
        # Rows imported before into another account: usually the wrong account was chosen for this file.
        counts = Counter()
        for row in job["rows"]:
            if (row.get("commit_status") or row["status"]) != "needs_review":
                continue
            for candidate in row.get("duplicate_candidates") or []:
                account = candidate.get("AccountNumber")
                if "import_identity_conflict" in candidate.get("reasons", []) and account is not None \
                        and str(account) != str(job.get("account_number")):
                    counts[str(account)] += 1
                    break
        return dict(counts)

    def summary(self, job):
        rows = job["rows"]
        # Where the categories of expenses still to import came from: memory (or the ShopName table), the model, the
        # user, or nowhere yet. Shown above the preview, so a large file says how much is left to classify by hand.
        flags = [self._open_expense(row) for row in rows]
        sources = Counter(row.get("classification_source") for row, (expense, unclassified) in zip(rows, flags)
                          if expense and not unclassified)
        return {**{k: v for k, v in job.items() if k not in {"rows", "source_path", "commits", "classification_cache"}},
                "counts": dict(Counter(row.get("commit_status") or row["status"] for row in rows)), "total": len(rows),
                "groups": {group: 0 for group in self.GROUPS} | dict(Counter(self.row_group(row) for row in rows)),
                "previous_accounts": self.previous_accounts(job),
                "classification": {"memory": sources["classification_memory"] + sources["shop_name_table"],
                                   "model": sources["llm"], "user": sources["user_confirmed"],
                                   "open": sum(unclassified for _, unclassified in flags)}}

    def get(self, ident):
        return self.summary(self.load(ident))

    def rows(self, ident, status=None, offset=0, limit=50):
        rows = self.load(ident)["rows"]
        if status == "unclassified":
            # Not a row status: expenses still lacking a category, the preview's "to classify" filter.
            rows = [row for row in rows if self._open_expense(row)[1]]
        elif status:
            rows = [row for row in rows if (row.get("commit_status") or row["status"]) == status]
        return {"rows": rows[offset:offset + limit], "total": len(rows), "offset": offset,
                "next_offset": offset + limit if offset + limit < len(rows) else None}

    def edit_rows(self, ident, revision, changes):
        with self.lock(ident):
            job = self.load(ident)
            if job["state"] not in {"ready", "partial_failed"} or job["revision"] != revision:
                raise ImportProblem("Preview is stale or not editable", 409)
            indexed = {row["row_id"]: row for row in job["rows"]}
            allowed = {"shop_name", "main_category", "sub_category", "detail_category"}
            with self.connect() as conn:
                for change in changes:
                    row = indexed.get(change["row_id"])
                    if not row or row.get("commit_status") in {"imported", "skipped"} or row["status"] not in {"new", "suspected_duplicate"}:
                        raise ImportProblem("Row cannot be edited", 409)
                    values = change["values"]
                    if not set(values) <= allowed or any(not isinstance(v, str) or len(v) > 300 for v in values.values()):
                        raise ImportProblem("Only merchant/category corrections are supported")
                    row["normalized"].update(values)
                    row["issues"] = [issue for issue in row["issues"] if issue not in {"classification_required", "classification_uncertain"}]
                    n = row["normalized"]
                    if not n["shop_name"].strip() or (n.get("transaction_kind") != "income" and
                            (not n["main_category"].strip() or not n["sub_category"].strip())):
                        raise ImportProblem("Merchant and categories cannot be empty")
                    row["status"], row["duplicate_candidates"] = find_duplicate_transaction(conn, row)
                    row["classification_source"] = "user_confirmed"
                    row.pop("commit_status", None)
                job["revision"] += 1
                self.save(job)
            return self.summary(job)

    def commit(self, ident, revision, choices, key):
        with self.lock(ident):
            job = self.load(ident)
            # Preserve old expense-only commit fingerprints when the new flag is absent/false.
            choices = [{k: v for k, v in choice.items() if k != "confirm_external_income" or v}
                       for choice in choices]
            digest = hashlib.sha256(json.dumps({"revision": revision, "rows": choices}, sort_keys=True).encode()).hexdigest()
            if key in job["commits"]:
                entry = job["commits"][key]
                if entry["hash"] != digest:
                    raise ImportProblem("Idempotency-Key already used with different choices", 409)
                if entry["state"] == "completed":
                    return self.summary(job)
            elif job["revision"] != revision or job["state"] not in {"ready", "partial_failed"}:
                raise ImportProblem("Preview is stale or task is not ready", 409)
            indexed = {row["row_id"]: row for row in job["rows"]}
            if len({x["row_id"] for x in choices}) != len(choices):
                raise ImportProblem("Duplicate row selections")
            for choice in choices:
                row = indexed.get(choice["row_id"])
                if row is None:
                    raise ImportProblem("Unknown row ID")
                if choice["action"] == "import" and not row.get("commit_status") == "imported":
                    n = row["normalized"]
                    if row["status"] not in {"new", "suspected_duplicate"} or (n.get("transaction_kind") != "income"
                            and (not n.get("main_category") or not n.get("sub_category"))):
                        raise ImportProblem("Resolve invalid/refund/unclassified rows before importing", 409)
                    if row.get("income_confirmation_required") and not choice.get("confirm_external_income"):
                        raise ImportProblem("Confirm this receipt came from another person, not your own account", 409)
                    if row["status"] == "suspected_duplicate" and not choice.get("confirm_distinct"):
                        raise ImportProblem("Explicitly confirm a suspected duplicate is a distinct purchase", 409)
            if not job.get("backup_path"):
                backup_path = self.folder(ident) / "before_commit.sqlite3"
                with self.connect() as source, closing(sqlite3.connect(backup_path)) as backup:
                    source.backup(backup)
                job["backup_path"] = str(backup_path)
            job["commits"][key] = {"hash": digest, "state": "running"}
            job["state"] = "importing"
            self.save(job)
            # Rows are written SAVE_EVERY at a time: one ledger transaction (one disk sync) per group, a savepoint per row
            # so a failing row is rolled back alone, and row states recorded only after the group commits, so the job
            # never claims a row the ledger does not have. run_operation is idempotent per row, so a retry after a
            # crash re-runs a group safely.
            pending = [(indexed[choice["row_id"]], choice) for choice in choices
                       if indexed[choice["row_id"]].get("commit_status") not in {"imported", "skipped"}]
            for start in range(0, len(pending), self.SAVE_EVERY):
                group, staged = pending[start:start + self.SAVE_EVERY], []
                try:
                    with self.connect() as conn:
                        conn.execute("BEGIN IMMEDIATE")
                        for row, choice in group:
                            if choice["action"] == "skip":
                                staged.append((row, {"commit_status": "skipped"}))
                                continue
                            conn.execute("SAVEPOINT import_row")
                            try:
                                staged.append((row, self._write_row(conn, job, ident, row, choice)))
                                conn.execute("RELEASE import_row")
                            except Exception as error:
                                conn.execute("ROLLBACK TO import_row")
                                conn.execute("RELEASE import_row")
                                staged.append((row, {"commit_status": "failed", "commit_error": str(error)[:500]}))
                except Exception as error:
                    # The group's transaction did not commit, so none of its rows reached the ledger.
                    staged = [(row, {"commit_status": "failed", "commit_error": str(error)[:500]}) for row, _ in group]
                for row, updates in staged:
                    row.update(updates)
                    if updates["commit_status"] != "failed":
                        row.pop("commit_error", None)
                self.save(job)
            job["commits"][key]["state"] = "completed"
            # Expenses written with a model or user category teach classification memory afterwards (teach_memory);
            # ones answered by memory or the ShopName table are already known.
            for row in job["rows"]:
                if (row.get("commit_status") == "imported" and "memory" not in row
                        and row.get("classification_source") in {"llm", "user_confirmed"}
                        and row["normalized"].get("transaction_kind") != "income"):
                    row["memory"] = "pending"
            remaining = [row for row in job["rows"] if row["status"] in {"new", "suspected_duplicate"} and row.get("commit_status") not in {"imported", "skipped"}]
            job["state"] = "partial_failed" if remaining or any(row.get("commit_status") in {"failed", "needs_review"} for row in job["rows"]) else "completed"
            job["revision"] += 1
            self.save(job)
            return self.summary(job)

    def _write_row(self, conn, job, ident, row, choice):
        """Write one previewed row inside the caller's transaction and return the row's new fields."""
        state, candidates = find_duplicate_transaction(conn, row)
        known = {(c.get("DealingNumber"), c.get("income_id")) for c in row["duplicate_candidates"]}
        current = {(c.get("DealingNumber"), c.get("income_id")) for c in candidates}
        if state == "duplicate":
            source_id = source_id_for_import(ident, row["row_id"])
            if candidates and candidates[0].get("DealingNumber"):
                attach_source(conn, source_id, candidates[0]["DealingNumber"], "statement", {
                    "filename": job.get("filename"),
                    "row_id": row.get("row_id"),
                })
                merge_merchant_if_generic(conn, candidates[0]["DealingNumber"], row["normalized"].get("shop_name"))
            return {"commit_status": "skipped", "duplicate_candidates": candidates}
        if state == "needs_review" or (state == "suspected_duplicate" and (not choice.get("confirm_distinct") or not current <= known)):
            return {"commit_status": "needs_review", "duplicate_candidates": candidates, "status": state}
        n = row["normalized"]
        account = assets.get_account(conn, n["AccountNumber"])
        if assets.currency(account["CurrencyType"]) != assets.currency(n["CurrencyTypeNumber"]):
            raise ImportProblem("Account currency changed since preview", 409)
        if n.get("transaction_kind") == "income":
            payload = {"account_number": n["AccountNumber"], "source_name": n["shop_name"],
                "amount": n["price"], "currency": assets.currency(n["CurrencyTypeNumber"]),
                "effective_at": n["isotime"], "note": n["product_name"],
                "provider": n["provider"], "transaction_id": n["transaction_id"],
                "merchant_order_id": n["merchant_order_id"], "payment_method": n["payment_method"],
                "main_category": n["main_category"], "sub_category": n["sub_category"],
                "detail_category": n["detail_category"], "filename": job["filename"],
                "import_id": ident, "row_number": row["row_number"],
                "external_income_confirmed": bool(choice.get("confirm_external_income"))}
            result = assets.run_operation(conn, row["operation_id"], "income.create", payload,
                lambda operation_id: assets.record_income(conn, payload, operation_id))
            return {"commit_status": "imported", "income_id": result["income_id"],
                    "balance_sync": {"status": "linked", "change_id": result["change_id"]}}
        payload = DBmanage.dealing_payload(n["shop_name"], n["price"], n["CurrencyTypeNumber"], n["isotime"],
            n["main_category"], n["sub_category"], n["AccountNumber"], n.get("detail_category", ""), None, None)
        result = assets.run_operation(conn, row["operation_id"], "dealing.create", payload,
            lambda operation_id: DBmanage.insert_dealing(conn, payload, operation_id))
        attach_source(conn, source_id_for_import(ident, row["row_id"]), result["DealingNumber"], "statement", {
            "filename": job.get("filename"),
            "row_id": row.get("row_id"),
        })
        return {"commit_status": "imported", "DealingNumber": result["DealingNumber"], "balance_sync": result["balance_sync"]}

    def teach_memory(self, ident):
        """Hand rows committed with model or user categories to the Agent's classification memory, so the next import
        of those merchants is answered from memory. Best effort: the ledger is already written, and rows whose
        lesson failed are retried after the next commit of this import."""
        with self.lock(ident):
            job = self.load(ident)
            rows = [row for row in job["rows"] if row.get("memory") in {"pending", "failed"}
                    and (row["normalized"].get("raw_shop_name") or row["normalized"].get("shop_name"))]
        if not rows:
            return
        lessons = [{"raw_shop_name": (n.get("raw_shop_name") or n["shop_name"])[:300], "shop_name": n["shop_name"][:300],
                    "product_name": (n.get("product_name") or "")[:4000], "main_category": n["main_category"][:100],
                    "sub_category": n["sub_category"][:100], "detail_category": (n.get("detail_category") or "")[:100],
                    "amount": str(n.get("price") or "")[:40], "iso_time": str(n.get("isotime") or "")[:40],
                    "source_id": f"import:{ident}:{row['row_id']}"[:200]}
                   for row in rows for n in [row["normalized"]]]
        try:
            result, state = self.call_agent("/agent/import/remember", {"rows": lessons}), "taught"
        except Exception as error:
            result, state = {"error": str(error)[:200]}, "failed"
        with self.lock(ident):
            job = self.load(ident)
            sent = {row["row_id"] for row in rows}
            for row in job["rows"]:
                if row["row_id"] in sent:
                    row["memory"] = state
            job["memory_sync"] = {"state": state, **(result if isinstance(result, dict) else {})}
            self.save(job)
